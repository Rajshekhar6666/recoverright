from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parents[2] / "data" / "synthetic_docs" / "sample_knee.docx"

doc = Document()
doc.add_heading("DISCHARGE SUMMARY", 0)
doc.add_paragraph("FICTIONAL SAMPLE - FOR SOFTWARE TESTING ONLY - NOT CLINICAL CONTENT")
doc.add_paragraph("Patient: Jane Example (fictional)")

doc.add_heading("Procedure", 1)
doc.add_paragraph("Procedure: Total knee replacement, right knee.")

doc.add_heading("Medications", 1)
rows = [
    ("Medication", "Dose", "Frequency", "Duration"),
    ("Examplamol", "500 mg", "Every 6 hours", "7 days"),
    ("Testocillin", "250 mg", "Twice daily", "5 days"),
    ("Placebex", "1-2 tablets", "As needed", "Until review"),  # ambiguous dose trap
]
table = doc.add_table(rows=len(rows), cols=4)
table.style = "Table Grid"
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        table.cell(r, c).text = val

doc.add_heading("Wound care", 1)
doc.add_paragraph("Keep the dressing clean and dry.", style="List Bullet")
doc.add_paragraph("Do not soak the wound in water until the stitches are removed.", style="List Bullet")

doc.add_heading("Activity", 1)
doc.add_paragraph("Walk with the frame for short distances several times a day.", style="List Bullet")
doc.add_paragraph("Do not drive until reviewed by the surgeon.", style="List Bullet")

doc.add_heading("Diet", 1)
doc.add_paragraph("Eat a normal balanced diet and drink plenty of fluids.")

doc.add_heading("Follow-up", 1)
doc.add_paragraph("Wound check at the surgical clinic. Date to be confirmed.")  # missing date trap
doc.add_paragraph("Physiotherapy review in 2 weeks.")

doc.add_heading("When to get help", 1)
doc.add_paragraph("Contact the ward or emergency services if you notice heavy bleeding from the wound, "
                  "new chest pain, or redness and swelling that is getting worse.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print("saved", OUT)