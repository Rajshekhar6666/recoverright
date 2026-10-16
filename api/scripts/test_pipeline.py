"""Small self-checking tests. Run:  python -m scripts.test_pipeline"""
import tempfile
from pathlib import Path

from docx import Document

from app.pipeline.ground import validate_plan
from app.pipeline.parse import parse_document
from app.pipeline.tables import table_medications
from app.pipeline.units import segment, split_sentences
from app.schemas.plan import ExtractedPlan, FollowUp, Medication, WoundCare

ROOT = Path(__file__).resolve().parents[2]
passed = 0


def check(name, cond):
    global passed
    print(("PASS  " if cond else "FAIL  ") + name)
    if not cond:
        raise SystemExit(1)
    passed += 1


# ---------- parser: duplicate cells, empty cells, merged cells ----------
tmp = Path(tempfile.mkdtemp())
d = Document()
d.add_heading("Medications", 1)
t = d.add_table(rows=4, cols=4)
data = [("Medication", "Dose", "Frequency", "Duration"),
        ("DrugOne", "1", "1", "5 days"),            # adjacent equal cells
        ("DrugTwo", "5 days", "", "5 days"),         # empty frequency, repeated text
        ("DrugThree", "10 mg", "Daily", "3 days")]
for r, row in enumerate(data):
    for c, v in enumerate(row):
        t.cell(r, c).text = v
t.cell(3, 1).merge(t.cell(3, 2))                     # merged cell spanning two columns
d.save(str(tmp / "t.docx"))
doc = parse_document(tmp / "t.docx")
rows = [b for b in doc.blocks if b.type == "table_row"]
check("adjacent equal cells are kept ('1' next to '1')", rows[1].cells == ["DrugOne", "1", "1", "5 days"])
check("empty cells keep their column", rows[2].cells == ["DrugTwo", "5 days", "", "5 days"])
check("merged cell appears once", len(rows[3].cells) == 3)
meds = {m.name: m for m in table_medications(doc)}
check("medication read by column header", meds["DrugOne"].dose == "1" and meds["DrugOne"].frequency == "1")
check("empty frequency becomes None, not shifted", meds["DrugTwo"].frequency is None and meds["DrugTwo"].duration == "5 days")
check("ragged (merged) row is skipped, not guessed", "DrugThree" not in meds)

# ---------- units ----------
check("abbreviation does not split", split_sentences("Dr. Smith will review. Come back soon.") == ["Dr. Smith will review.", "Come back soon."])
check("Hindi danda splits", len(split_sentences("घाव को साफ रखें। भारी वजन न उठाएं।")) == 2)
check("Chinese full stop splits", len(split_sentences("保持伤口干燥。不要开车。")) == 2)

# ---------- validator on the real sample ----------
sample = parse_document(ROOT / "data" / "synthetic_docs" / "sample_knee.docx")
row = next(b for b in sample.blocks if "Examplamol" in b.text)
row_q = row.text
units = segment(sample)
check("units carry table headers", any(u.col_header == "dose" and u.text == "500 mg" for u in units))
check("units carry the section heading", any(u.section == "Medications" for u in units))


def plan_with(**med):
    base = dict(name="Examplamol", dose="500 mg", frequency="Every 6 hours", duration="7 days",
                source_quote=row_q, block_id=row.block_id)
    base.update(med)
    return ExtractedPlan(medications=[Medication(**base)])


good = validate_plan(plan_with(), sample)[0]
check("correct medication passes", good.verdict == "pass")
bad = validate_plan(plan_with(duration="6 days"), sample)[0]
check("duration '6 days' assembled from '6 hours' + '7 days' FAILS", bad.verdict == "fail")
check("...with a helpful explanation", any("different cells" in i for i in bad.issues))
check("wrong dose number fails", validate_plan(plan_with(dose="50 mg"), sample)[0].verdict == "fail")
check("swapped columns fail", validate_plan(plan_with(dose="7 days"), sample)[0].verdict == "fail")

wc = next(b for b in sample.blocks if b.text.startswith("Keep the dressing"))
ok = ExtractedPlan(wound_care=[WoundCare(instruction="Keep the dressing clean and dry", source_quote=wc.text, block_id=wc.block_id)])
check("trailing period is not an error", validate_plan(ok, sample)[0].verdict == "pass")
changed = ExtractedPlan(wound_care=[WoundCare(instruction="Keep the dressing dry for 3 days", source_quote=wc.text, block_id=wc.block_id)])
check("invented number in free text fails", validate_plan(changed, sample)[0].verdict == "fail")

fu = next(b for b in sample.blocks if "Wound check" in b.text)
ph = ExtractedPlan(follow_ups=[FollowUp(what="Wound check", when="Date to be confirmed", source_quote=fu.text, block_id=fu.block_id)])
check("placeholder date fails", validate_plan(ph, sample)[0].verdict == "fail")
nb = ExtractedPlan(wound_care=[WoundCare(instruction="Keep the dressing clean and dry", source_quote=wc.text, block_id="b999")])
check("non-existent block fails", validate_plan(nb, sample)[0].verdict == "fail")

print(f"\nall {passed} checks passed")
