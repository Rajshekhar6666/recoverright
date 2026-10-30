import re
from pathlib import Path
from typing import Iterator, Optional

import pdfplumber
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.schemas.document import Block, SourceDocument
from app.schemas.events import Stage
from app.services.event_bus import bus


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\xa0", " ")).strip()


def _raw(type_, text, page=None, table_id=None, row_index=None, cells=None) -> dict:
    return {"type": type_, "text": text, "page": page,
            "table_id": table_id, "row_index": row_index, "cells": cells}


def _iter_docx(path: Path) -> Iterator[dict]:
    doc = Document(str(path))
    table_no = 0
    for child in doc.element.body.iterchildren():  # paragraphs and tables, in document order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            ppr = p._p.pPr
            if style.startswith("Heading") or style == "Title":
                yield _raw("heading", text)
            elif "List" in style or (ppr is not None and ppr.numPr is not None):
                yield _raw("list_item", text)
            else:
                yield _raw("paragraph", text)
        elif tag == "tbl":
            table_no += 1
            table = Table(child, doc)
            for r, row in enumerate(table.rows):
                cells: list[str] = []
                seen_tcs = []  # keep references alive so identity checks are reliable
                for cell in row.cells:
                    # a merged cell is returned once per column it spans; skip repeats by identity,
                    # never by text, so "1" next to "1" or "5 days | 5 days" survive
                    if any(cell._tc is tc for tc in seen_tcs):
                        continue
                    seen_tcs.append(cell._tc)
                    cells.append(_clean(cell.text))
                if any(cells):
                    yield _raw("table_row", " | ".join(cells), None, f"t{table_no}", r, cells)


def _pdf_line(text: str, page_no: int, chars: list, body_size: float) -> dict:
    if re.match(r"^([-*\u2022]|\d+[.)])\s+", text):
        return _raw("list_item", text, page_no)
    sizes = [c.get("size", body_size) for c in chars] or [body_size]
    bold = sum("bold" in str(c.get("fontname", "")).lower() for c in chars) / max(len(chars), 1) >= 0.8
    big = (sum(sizes) / len(sizes)) >= body_size * 1.15
    if len(text) <= 80 and not re.search(r"[.!?\u0964\u3002]$", text) and (bold or big or (text.isupper() and len(text) <= 60)):
        return _raw("heading", text, page_no)
    return _raw("paragraph", text, page_no)


_SENTENCE_END = re.compile(r"[.!?:;\u0964\u3002]$")


def _iter_text(path: Path) -> Iterator[dict]:
    content = ""
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            content = path.read_text(encoding=enc)
            break
        except Exception:
            continue
    if not content:
        return
    for line in content.splitlines():
        line = _clean(line)
        if not line:
            continue
        if re.match(r"^([-*\u2022]|\d+[.)])\s+", line):
            yield _raw("list_item", line)
        elif (line.endswith(":") and len(line) <= 80) or (line.isupper() and len(line) <= 60) or line.startswith("#"):
            yield _raw("heading", line.lstrip("#").strip())
        elif "|" in line and line.count("|") >= 2:
            cells = [_clean(c) for c in line.split("|") if _clean(c)]
            if cells:
                yield _raw("table_row", " | ".join(cells), None, None, None, cells)
        else:
            yield _raw("paragraph", line)


def _iter_pdf(path: Path) -> Iterator[dict]:
    found_any = False
    with pdfplumber.open(str(path)) as pdf:
        table_no = 0
        for page_no, page in enumerate(pdf.pages, 1):
            sizes = sorted(c["size"] for c in page.chars) or [10.0]
            body_size = sizes[len(sizes) // 2]  # median character size = body text
            entries: list[tuple[float, dict]] = []
            tables = []
            try:
                tables = page.find_tables() or []
            except Exception:
                pass
            region = page
            for t in tables:
                try:
                    region = region.outside_bbox(t.bbox)  # text outside tables, so nothing is read twice
                except Exception:
                    pass
            prev = None  # last paragraph/list entry that a wrapped line may continue
            try:
                for line in region.extract_text_lines():
                    text = _clean(line["text"])
                    if not text:
                        continue
                    raw = _pdf_line(text, page_no, line.get("chars", []), body_size)
                    height = max(line["bottom"] - line["top"], 1.0)
                    if (prev is not None and raw["type"] == "paragraph"
                            and prev["raw"]["type"] in ("paragraph", "list_item")
                            and not _SENTENCE_END.search(prev["raw"]["text"])
                            and (line["top"] - prev["bottom"]) < 0.35 * height):
                        prev["raw"]["text"] += " " + text  # a wrapped line continues its paragraph
                        prev["bottom"] = line["bottom"]
                        continue
                    prev = {"raw": raw, "bottom": line["bottom"]}
                    entries.append((line["top"], raw))
            except Exception:
                pass

            for t in tables:
                try:
                    table_no += 1
                    for r, row in enumerate(t.extract() or []):
                        cells = [_clean(c or "") for c in row]
                        if any(cells):
                            entries.append((t.bbox[1] + r * 0.01,
                                            _raw("table_row", " | ".join(cells), page_no, f"t{table_no}", r, cells)))
                except Exception:
                    pass

            # Fallback A: if no entries extracted from text lines, extract page text directly
            if not entries:
                try:
                    raw_page_text = page.extract_text() or ""
                    for l_idx, line in enumerate(raw_page_text.splitlines()):
                        cleaned = _clean(line)
                        if cleaned:
                            entries.append((float(l_idx), _pdf_line(cleaned, page_no, [], body_size)))
                except Exception:
                    pass

            for _, raw in sorted(entries, key=lambda x: x[0]):
                found_any = True
                yield raw

    # Fallback B: if pdfplumber found nothing at all (e.g. font dictionary or layout issue), try pypdfium2
    if not found_any:
        try:
            import pypdfium2 as pdfium
            pdf_doc = pdfium.PdfDocument(str(path))
            for p_no in range(len(pdf_doc)):
                page = pdf_doc.get_page(p_no)
                textpage = page.get_textpage()
                raw_text = textpage.get_text_range()
                for line in raw_text.splitlines():
                    cleaned = _clean(line)
                    if cleaned:
                        if re.match(r"^([-*\u2022]|\d+[.)])\s+", cleaned):
                            yield _raw("list_item", cleaned, p_no + 1)
                        elif (cleaned.endswith(":") and len(cleaned) <= 80) or (cleaned.isupper() and len(cleaned) <= 60):
                            yield _raw("heading", cleaned, p_no + 1)
                        else:
                            yield _raw("paragraph", cleaned, p_no + 1)
        except Exception:
            pass


def _ocr_image(path: Path) -> list[dict]:
    """OCR fallback for image uploads (.png, .jpg, etc.) using pytesseract if available."""
    entries: list[dict] = []
    try:
        import pytesseract
        from PIL import Image
        import shutil

        if not shutil.which("tesseract"):
            for cand in [
                Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
                Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
                Path(r"/usr/bin/tesseract"),
                Path(r"/usr/local/bin/tesseract"),
            ]:
                if cand.exists():
                    pytesseract.pytesseract.tesseract_cmd = str(cand)
                    break

        img = Image.open(str(path))
        text = pytesseract.image_to_string(img)
        for line in text.splitlines():
            cleaned = _clean(line)
            if not cleaned:
                continue
            if re.match(r"^([-*\u2022]|\d+[.)])\s+", cleaned):
                entries.append(_raw("list_item", cleaned))
            elif len(cleaned) <= 60 and not re.search(r"[.!?\u0964\u3002]$", cleaned):
                entries.append(_raw("heading", cleaned))
            else:
                entries.append(_raw("paragraph", cleaned))
    except Exception:
        pass
    return entries


def _ocr_pdf(path: Path) -> list[dict]:
    """OCR fallback for image-only/scanned PDFs using pdfplumber page rasterization and pytesseract."""
    entries: list[dict] = []
    try:
        import pytesseract
        import shutil

        # Discover tesseract binary on host or standard installation paths
        if not shutil.which("tesseract"):
            for cand in [
                Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
                Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
                Path(r"/usr/bin/tesseract"),
                Path(r"/usr/local/bin/tesseract"),
            ]:
                if cand.exists():
                    pytesseract.pytesseract.tesseract_cmd = str(cand)
                    break

        with pdfplumber.open(str(path)) as pdf:
            for page_no, page in enumerate(pdf.pages, 1):
                try:
                    img = page.to_image(resolution=150).original
                    text = pytesseract.image_to_string(img)
                    for line in text.splitlines():
                        cleaned = _clean(line)
                        if not cleaned:
                            continue
                        if re.match(r"^([-*\u2022]|\d+[.)])\s+", cleaned):
                            entries.append(_raw("list_item", cleaned, page_no))
                        elif len(cleaned) <= 60 and not re.search(r"[.!?\u0964\u3002]$", cleaned):
                            entries.append(_raw("heading", cleaned, page_no))
                        else:
                            entries.append(_raw("paragraph", cleaned, page_no))
                except Exception:
                    continue
    except Exception:
        pass
    return entries


def parse_document(path: str | Path, run_id: Optional[str] = None) -> SourceDocument:
    path = Path(path)

    def emit(type_: str, payload: dict | None = None) -> None:
        if run_id:
            bus.emit(run_id, Stage.PARSE, type_, payload)

    emit("parse.started", {"filename": path.name})
    ext = path.suffix.lower()
    if ext == ".docx":
        raw = list(_iter_docx(path))
    elif ext == ".pdf":
        raw = list(_iter_pdf(path))
        if not raw:
            emit("parse.ocr", {"message": "No embedded digital text found; invoking OCR engine on scanned PDF."})
            raw = _ocr_pdf(path)
    elif ext in {".txt", ".text", ".md", ".rtf", ".csv"}:
        raw = list(_iter_text(path))
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        emit("parse.ocr", {"message": f"Extracting text from {ext} image via OCR."})
        raw = _ocr_image(path)
    else:
        emit("parse.failed", {"reason": f"unsupported file type {ext}"})
        raise ValueError(f"Unsupported file type: {ext} (supported: .docx, .pdf, .txt, .md, .png, .jpg)")

    if not raw:
        emit("parse.failed", {"reason": "no text found in document"})
        raise ValueError(
            "No extractable text found in this document. If this is a photo or scanned PDF without embedded text, "
            "please paste the clinical notes directly or provide a digital document."
        )

    blocks: list[Block] = []
    for i, r in enumerate(raw, 1):
        block = Block(block_id=f"b{i}", type=r["type"], page=r["page"], text=r["text"],
                      table_id=r["table_id"], row_index=r["row_index"], cells=r["cells"])
        blocks.append(block)
        emit("parse.block", {"block_id": block.block_id, "type": block.type, "preview": block.text[:80]})

    emit("parse.done", {"blocks": len(blocks)})
    return SourceDocument(filename=path.name, blocks=blocks)
